#!/usr/bin/env python3
"""Dump Order-of-Service / notices emails as text, keeping run formatting.

Usage:
    .venv/bin/python3 .claude/skills/order-of-service/scripts/extract_oos.py \
        OUT_DIR path/to/one.eml [path/to/two.eml ...]

For every email: prints headers, the plain-text body and the stripped HTML
body (for line layout and link targets), recurses into attached emails, saves every
attachment into OUT_DIR, and prints the text of .docx / .pdf attachments.

Formatting is made visible because responsive readings are often marked
only by type style (light vs bold, or coloured text):
    **bold**   _italic_   {color:#RRGGBB}text{/color}   {highlight:yellow}text{/highlight}
Black / automatic colour is not marked.
"""
import sys
from email import policy
from email.parser import BytesParser
from pathlib import Path

from bs4 import BeautifulSoup

PLAIN_COLOURS = {None, "000000", "auto"}


def docx_effective(run, para, attr):
    """Resolve a run's bold/italic through run -> char style -> paragraph style chain."""
    value = getattr(run.font, attr)
    if value is not None:
        return value
    style = run.style
    while style is not None:
        if getattr(style.font, attr) is not None:
            return getattr(style.font, attr)
        style = style.base_style
    style = para.style
    while style is not None:
        if getattr(style.font, attr) is not None:
            return getattr(style.font, attr)
        style = style.base_style
    return False


def docx_colour(run, para):
    for font in (run.font, getattr(run.style, "font", None), para.style.font):
        if font is not None and font.color is not None and font.color.type is not None:
            return str(font.color.rgb) if font.color.rgb is not None else None
    return None


def format_docx_paragraph(para):
    out = []
    for run in para.runs:
        text = run.text
        if not text.strip():
            out.append(text)
            continue
        if docx_effective(run, para, "italic"):
            text = f"_{text}_"
        if docx_effective(run, para, "bold"):
            text = f"**{text}**"
        colour = docx_colour(run, para)
        if colour not in PLAIN_COLOURS:
            text = f"{{color:#{colour}}}{text}{{/color}}"
        if run.font.highlight_color is not None:
            text = f"{{highlight:{run.font.highlight_color}}}{text}{{/highlight}}"
        out.append(text)
    # Merge adjacent identical markers so "**a****b**" reads as "**ab**".
    return "".join(out).replace("****", "")


def dump_docx(path):
    import docx

    document = docx.Document(path)
    for para in document.paragraphs:
        line = format_docx_paragraph(para)
        if line.strip():
            print(f"[{para.style.name}] {line}")
    for t_index, table in enumerate(document.tables, 1):
        print(f"--- table {t_index} ---")
        for row in table.rows:
            print(" | ".join(" / ".join(format_docx_paragraph(p) for p in cell.paragraphs) for cell in row.cells))
    images = [rel.target_ref for rel in document.part.rels.values() if "image" in rel.reltype]
    links = [rel.target_ref for rel in document.part.rels.values() if "hyperlink" in rel.reltype]
    print(f"--- embedded images: {images or 'none'}")
    print(f"--- hyperlinks: {links or 'none'}")


def dump_pdf(path):
    import pymupdf

    with pymupdf.open(path) as pdf:
        for page_no, page in enumerate(pdf, 1):
            print(f"--- page {page_no} ---")
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    parts = []
                    for span in line["spans"]:
                        text = span["text"]
                        if not text.strip():
                            parts.append(text)
                            continue
                        if span["flags"] & 2:
                            text = f"_{text}_"
                        if span["flags"] & 16 or "bold" in span["font"].lower():
                            text = f"**{text}**"
                        colour = f"{span['color']:06X}"
                        if colour != "000000":
                            text = f"{{color:#{colour}}}{text}{{/color}}"
                        parts.append(text)
                    if "".join(parts).strip():
                        print("".join(parts))
            links = [l["uri"] for l in page.get_links() if l.get("uri")]
            if links:
                print(f"--- page {page_no} links: {links}")


def dump_message(msg, out_dir, depth=0):
    bar = "=" * (5 + depth * 2)
    print(f"{bar} EMAIL (depth {depth})")
    for header in ("From", "To", "Cc", "Date", "Subject"):
        if msg[header]:
            print(f"{header}: {msg[header]}")

    plain = msg.get_body(preferencelist=("plain",))
    if plain is not None:
        print("--- body (text/plain) ---")
        print(plain.get_content())
    html = msg.get_body(preferencelist=("html",))
    if html is not None:
        # Printed even when a plain part exists: notices sheets often have a
        # plain part whose line breaks have collapsed, while the HTML keeps them.
        print("--- body (text/html, stripped: use for line layout, and hrefs) ---")
        soup = BeautifulSoup(html.get_content(), "html.parser")
        lines = (" ".join(line.split()) for line in soup.get_text("\n").splitlines())
        print("\n".join(line for line in lines if line))
        hrefs = [a["href"] for a in soup.find_all("a", href=True)]
        print(f"--- hrefs: {hrefs or 'none'}")

    for part in msg.iter_attachments():
        if part.get_content_type() == "message/rfc822":
            dump_message(part.get_content(), out_dir, depth + 1)
            continue
        name = part.get_filename() or f"attachment_{depth}.bin"
        target = out_dir / name
        target.write_bytes(part.get_payload(decode=True))
        print(f"{bar} ATTACHMENT: {name} ({part.get_content_type()}) -> {target}")
        suffix = target.suffix.lower()
        if suffix == ".docx":
            dump_docx(target)
        elif suffix == ".pdf":
            dump_pdf(target)
        elif suffix == ".eml":
            dump_message(BytesParser(policy=policy.default).parse(target.open("rb")), out_dir, depth + 1)
        else:
            print("(not parsed by this script -- handle per the kindle-doc-prep email rules)")


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    out_dir = Path(sys.argv[1])
    out_dir.mkdir(parents=True, exist_ok=True)
    for eml in sys.argv[2:]:
        with open(eml, "rb") as handle:
            dump_message(BytesParser(policy=policy.default).parse(handle), out_dir)


if __name__ == "__main__":
    main()
